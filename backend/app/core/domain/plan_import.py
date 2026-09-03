"""碳循环计划表的解析。

用户手上的计划多半是一份 Excel——健身博主发的模板、教练给的表格、自己拿
别人的表改的。让他对着表格一天天口述给助理是荒谬的：这份文件里有 28 天
× 6 个字段，任何人都会在第三天放弃。

## 为什么是纯函数、不碰数据库

解析是**格式问题**，落库是**业务问题**。混在一起时，"这个表能不能读懂"这件事
就只能连着数据库一起测，而格式的边界情况（少一列、多一列、单位写成中文）恰恰
是最需要密集测试的部分。所以这里只做 bytes → 结构化数据，一个 IO 都不碰。

## 真实文件是脏的

拿到的样例（凯圣王碳循环计划_91kg.xlsx）第 1 周有 **9 列**而其余三周是 8 列，
多出来的那列夹在"一"和"二"之间，值是 138.4 / 111.2 / 53.4 / 1491——像是有人
在旁边试算了一版没删掉。如果按"第 2..8 列就是周一到周日"硬读，第 1 周会整周
错位一天，而且不会报错：数字都在合理范围内，导入完看起来一切正常。

所以列位置一律**由表头的星期字符决定**，不靠序号。认不出的列直接丢掉。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

# 星期表头的所有写法。真实模板里这一格可能是"一""周一""星期一""Mon"，
# 甚至带空格。全部归一到 1-7。
_WEEKDAY_ALIASES: dict[str, int] = {}
for _index, _names in enumerate(
    [
        ("一", "周一", "星期一", "mon", "monday"),
        ("二", "周二", "星期二", "tue", "tuesday"),
        ("三", "周三", "星期三", "wed", "wednesday"),
        ("四", "周四", "星期四", "thu", "thursday"),
        ("五", "周五", "星期五", "fri", "friday"),
        ("六", "周六", "星期六", "sat", "saturday"),
        ("日", "周日", "星期日", "天", "周天", "sun", "sunday"),
    ],
    start=1,
):
    for _name in _names:
        _WEEKDAY_ALIASES[_name] = _index

# 行标签 → 字段名。用"包含"匹配而不是相等：真实表头带单位，
# 写法五花八门（"碳水(g)"、"碳水化合物 (克)"、"碳水 g"）。
#
# 顺序按特异性从高到低排，第一个命中即返回。这不是风格问题——样例表里
# "缺口/盈余(kcal)" 同时含有"缺口"和"kcal"，如果 kcal 规则排在前面，
# 缺口那一行会被认成热量，把真正的热量值覆盖掉。实测踩过。
_ROW_PATTERNS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("balance_kcal", ("缺口", "盈余", "赤字", "差值")),
    ("day_type", ("碳日", "碳水日", "日型", "类型")),
    ("focus", ("训练部位", "训练", "部位", "计划")),
    ("carb_g", ("碳水", "碳水化合物")),
    ("protein_g", ("蛋白",)),
    ("fat_g", ("脂肪",)),
    ("kcal", ("热量", "卡路里", "能量", "kcal")),
)

# 碳日类型的归一化。存英文是为了和 plan_cut_phase 生成的 payload 对齐
# （那边用 high/medium/low），否则同一张卡片要处理两套取值。
_DAY_TYPE_ALIASES: dict[str, str] = {
    "高": "high", "高碳": "high", "high": "high",
    "中": "medium", "中碳": "medium", "medium": "medium", "mid": "medium",
    "低": "low", "低碳": "low", "low": "low",
    "休息": "rest", "rest": "rest",
}

# 基础信息表的键映射。只挑能对上 profile 白名单的，其余留在 raw 里。
_PROFILE_KEYS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("weight_kg", ("当前体重", "体重")),
    ("target_kg", ("目标体重",)),
    ("height_cm", ("身高",)),
    ("age", ("年龄",)),
    ("sex", ("性别",)),
)

_SHEET_WEEK_RE = re.compile(r"第\s*(\d+)\s*周|week\s*(\d+)", re.IGNORECASE)

# 单个数值的合理区间 (下界, 上界)。超出的当脏数据丢掉而不是原样导入——
# 一份写着每日 30000 kcal 的计划进了库，后面每一次"今天该吃多少"都会答错。
#
# balance_kcal 的下界是负的：它是热量缺口，减脂计划里本来就该是负数。
# 一开始统一按 0 起判，结果把样例文件里全部 28 天的缺口都当成越界丢掉了。
_LIMITS: dict[str, tuple[float, float]] = {
    "carb_g": (0.0, 2000.0),
    "protein_g": (0.0, 1000.0),
    "fat_g": (0.0, 1000.0),
    "kcal": (0.0, 20000.0),
    "balance_kcal": (-20000.0, 20000.0),
}


class PlanParseError(ValueError):
    """解析失败。消息直接给用户看，所以必须说清楚是哪里的问题。"""


@dataclass
class ParsedDay:
    day: int
    day_type: str = "medium"
    focus: str = ""
    carb_g: float | None = None
    protein_g: float | None = None
    fat_g: float | None = None
    kcal: float | None = None
    balance_kcal: float | None = None

    def to_payload(self) -> dict:
        # 只输出有值的字段。None 落进 JSONB 后前端还要逐个判空，
        # 而"这一天没写脂肪"和"这一天脂肪是 0"是两件事。
        data: dict = {"day": self.day, "day_type": self.day_type}
        if self.focus:
            data["focus"] = self.focus
        for name in ("carb_g", "protein_g", "fat_g", "kcal", "balance_kcal"):
            value = getattr(self, name)
            if value is not None:
                data[name] = value
        return data


@dataclass
class ParsedWeek:
    week: int
    days: list[ParsedDay] = field(default_factory=list)

    def to_payload(self) -> dict:
        return {"week": self.week, "days": [d.to_payload() for d in self.days]}


@dataclass
class ParsedPlan:
    weeks: list[ParsedWeek]
    profile: dict = field(default_factory=dict)
    basics: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    @property
    def day_count(self) -> int:
        return sum(len(week.days) for week in self.weeks)

    def to_payload(self) -> dict:
        """转成 plans.payload。

        刻意与 plan_cut_phase 生成的形状保持兼容（都有 weeks[].days[]、
        都用 high/medium/low），这样同一个卡片组件和同一个查询工具能同时
        处理"助理生成的"和"用户导入的"计划。
        """
        return {
            "source": "import",
            "carb_cycle": any(
                day.day_type in ("high", "low")
                for week in self.weeks for day in week.days
            ),
            "basics": self.basics,
            "weeks": [week.to_payload() for week in self.weeks],
        }


def _clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in ("nan", "none", "-", "—") else text


def _match_label(text: str, patterns: tuple[str, ...]) -> bool:
    return any(pattern in text for pattern in patterns)


def _to_number(value: object) -> float | None:
    """抽出数值。表格里的数字常带单位或区间，不能直接 float()。"""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _clean(value)
    if not text:
        return None
    # "1500-1600 kcal" 取第一个数；"约 140g" 也能拿到 140
    match = re.search(r"-?\d+(?:\.\d+)?", text.replace(",", ""))
    return float(match.group()) if match else None


def _row_field(label: str) -> str | None:
    for name, patterns in _ROW_PATTERNS:
        if _match_label(label, patterns):
            return name
    return None


def _week_number(sheet_title: str, fallback: int) -> int:
    match = _SHEET_WEEK_RE.search(sheet_title)
    if match:
        digits = match.group(1) or match.group(2)
        if digits and digits.isdigit():
            return int(digits)
    return fallback


def _weekday_columns(header: tuple) -> dict[int, int]:
    """表头 → {列索引: 星期几}。

    这是整个解析里最要紧的一步。样例文件第 1 周多了一列试算数据，按列序号
    读会整周错位一天且不报错——数字都在合理范围内，导入完看起来一切正常。
    所以只认表头里的星期字符，认不出的列丢掉。
    """
    mapping: dict[int, int] = {}
    for index, cell in enumerate(header):
        key = _clean(cell).lower().replace(" ", "")
        weekday = _WEEKDAY_ALIASES.get(key)
        # 同一个星期出现两次时保留第一个：重复列通常是右边那份草稿。
        if weekday and weekday not in mapping.values():
            mapping[index] = weekday
    return mapping


def _parse_week_sheet(
    sheet, fallback_week: int, warnings: list[str],
) -> ParsedWeek | None:
    rows = [row for row in sheet.iter_rows(values_only=True) if any(_clean(c) for c in row)]
    if not rows:
        return None

    columns = _weekday_columns(rows[0])
    if not columns:
        warnings.append(f"工作表「{sheet.title}」没有识别到星期表头，已跳过")
        return None

    week = ParsedWeek(week=_week_number(sheet.title, fallback_week))
    days = {weekday: ParsedDay(day=weekday) for weekday in sorted(columns.values())}

    for row in rows[1:]:
        label = _clean(row[0] if row else "")
        if not label:
            continue
        field_name = _row_field(label)
        if field_name is None:
            continue
        for index, weekday in columns.items():
            if index >= len(row):
                continue
            raw = row[index]
            day = days[weekday]
            if field_name == "day_type":
                text = _clean(raw)
                day.day_type = _DAY_TYPE_ALIASES.get(text.lower(), text or "medium")
            elif field_name == "focus":
                day.focus = _clean(raw)
            else:
                number = _to_number(raw)
                bounds = _LIMITS.get(field_name)
                if number is not None and bounds is not None:
                    low, high = bounds
                    if not (low <= number <= high):
                        warnings.append(
                            f"{sheet.title} 周{weekday} 的{label}是 {number}，"
                            f"超出合理范围（{low:g}~{high:g}），已忽略",
                        )
                        number = None
                setattr(day, field_name, number)

    week.days = [days[key] for key in sorted(days)]
    return week


def _parse_basics(sheet) -> tuple[dict, dict]:
    """基础信息表 → (profile 可用字段, 原样保留的键值)。"""
    basics: dict = {}
    profile: dict = {}
    for row in sheet.iter_rows(values_only=True):
        if not row or len(row) < 2:
            continue
        label, raw = _clean(row[0]), row[1]
        if not label or _clean(raw) == "":
            continue
        basics[label] = raw if isinstance(raw, (int, float)) else _clean(raw)

        for name, patterns in _PROFILE_KEYS:
            if not _match_label(label, patterns) or name in profile:
                continue
            if name == "sex":
                text = _clean(raw)
                if text in ("男", "male", "m"):
                    profile["sex"] = "male"
                elif text in ("女", "female", "f"):
                    profile["sex"] = "female"
            else:
                number = _to_number(raw)
                if number is not None:
                    profile[name] = int(number) if name == "age" else number
            break
    return profile, basics


def parse_carb_cycle_workbook(data: bytes) -> ParsedPlan:
    """解析碳循环计划工作簿。

    只接受 bytes 而不是路径：调用方是 HTTP 上传，落地成临时文件纯属多余，
    还得管清理。
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - 依赖缺失时的明确报错
        raise PlanParseError("服务端缺少 Excel 解析依赖 openpyxl") from exc

    try:
        # data_only=True 取公式算出的值。样例表里"热量"是由宏量算出来的，
        # 不取值只会拿到 "=B4*4+B5*4+B6*9" 这样的字符串。
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl 的异常类型很杂
        raise PlanParseError(
            "这个文件读不出来，请确认是 .xlsx 格式（旧的 .xls 需要另存为 xlsx）",
        ) from exc

    try:
        warnings: list[str] = []
        profile: dict = {}
        basics: dict = {}
        weeks: list[ParsedWeek] = []

        for sheet in workbook.worksheets:
            title = _clean(sheet.title)
            # 基础信息表用"没有星期表头"来区分不可靠——空表也没有。
            # 按标题判断更直接，认错了也只是少填几个档案字段。
            if any(key in title for key in ("基础", "信息", "概览", "summary")):
                profile, basics = _parse_basics(sheet)
                continue
            week = _parse_week_sheet(sheet, len(weeks) + 1, warnings)
            if week and week.days:
                weeks.append(week)
    finally:
        # read_only 模式会持有文件句柄，不关会泄漏。
        workbook.close()

    if not weeks:
        raise PlanParseError(
            "没有在这个文件里找到周计划。每一周应该是一个工作表，"
            "第一行是星期（一/二/三…），左边一列是「碳日」「训练部位」「碳水(g)」这类标签。",
        )

    # 周号可能重复（两张表都叫"第1周"）或缺号。重排成连续序号，并保留告警——
    # 静默改号会让用户对不上自己的原表。
    seen: set[int] = set()
    for position, week in enumerate(weeks, start=1):
        if week.week in seen or week.week <= 0:
            warnings.append(f"第 {position} 张周表的周号是 {week.week}，与其他周重复，已改为 {position}")
            week.week = position
        seen.add(week.week)

    return ParsedPlan(weeks=weeks, profile=profile, basics=basics, warnings=warnings)
